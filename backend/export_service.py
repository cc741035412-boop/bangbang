"""从同一份观察记录数据生成 DOCX、PDF 和 Markdown。"""

from dataclasses import dataclass, field
from datetime import date, datetime
from html import escape
from io import BytesIO
from pathlib import Path
from typing import List, Optional, Tuple
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as ReportLabImage,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from time_utils import ensure_utc


KINDERGARTEN_TIMEZONE = ZoneInfo("Asia/Shanghai")
DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MEDIA_TYPE = "application/pdf"
MARKDOWN_MEDIA_TYPE = "text/markdown; charset=utf-8"
LEVEL_LABELS = {1: "初阶", 2: "中阶", 3: "高阶"}
FONT_NAME = "宋体"
PDF_FONT_NAME = "BangbangCJK"
PDF_FONT_CANDIDATES = (
    Path("/System/Library/Fonts/STHeiti Medium.ttc"),
    Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    Path("/usr/share/fonts/truetype/arphic/uming.ttc"),
)


@dataclass
class ExportChild:
    name: str
    birth_date: Optional[date] = None
    gender: Optional[str] = None


@dataclass
class ExportIndicator:
    code: str
    name: str
    level: int
    # 该层级的发展锚点描述（如"能持续进行身体操作（每次 30 分钟以上）…"）
    level_desc: str = ""
    # 白描原文回溯到的行为证据（如"白描原文：幼儿X持续搭高…"）
    evidence: str = ""


@dataclass
class ExportObservation:
    observation_id: int
    observed_at: datetime
    age_group: str
    area_name: str
    observer_name: str
    children: List[ExportChild] = field(default_factory=list)
    location: Optional[str] = None
    background_note: Optional[str] = None
    purpose: Optional[str] = None
    narrative: Optional[str] = None
    analysis: Optional[str] = None
    strategy: Optional[str] = None
    indicators: List[ExportIndicator] = field(default_factory=list)
    # 上传照片与视频代表帧，导出时统一作为现场影像展示。
    highlight_frames: List[bytes] = field(default_factory=list)


def kindergarten_datetime(value: datetime) -> datetime:
    """观察文书统一按幼儿园所在地 Asia/Shanghai 展示。"""
    return ensure_utc(value).astimezone(KINDERGARTEN_TIMEZONE)


def academic_year_and_term(value: datetime) -> Tuple[str, str]:
    """按园所当前规则计算学年度；其他园所规则不同可只替换此函数。"""
    local = kindergarten_datetime(value)
    if 9 <= local.month <= 12:
        return f"{local.year}-{local.year + 1}学年度", "第一学期"
    if local.month == 1:
        return f"{local.year - 1}-{local.year}学年度", "第一学期"
    return f"{local.year - 1}-{local.year}学年度", "第二学期"


def actual_age(birth_date: date, observed_at: datetime) -> int:
    observed_date = kindergarten_datetime(observed_at).date()
    return (
        observed_date.year
        - birth_date.year
        - ((observed_date.month, observed_date.day) < (birth_date.month, birth_date.day))
    )


def _collapse_same(values: List[str]) -> str:
    if not values:
        return ""
    return values[0] if all(value == values[0] for value in values) else "、".join(values)


def child_names(record: ExportObservation) -> str:
    return "、".join(child.name for child in record.children)


def child_ages(record: ExportObservation) -> str:
    age_group_labels = {"small": "小班", "middle": "中班", "large": "大班"}
    fallback = age_group_labels.get(record.age_group, record.age_group)
    values = [
        f"{actual_age(child.birth_date, record.observed_at)}岁"
        if child.birth_date
        else fallback
        for child in record.children
    ]
    return _collapse_same(values) if values else fallback


def child_genders(record: ExportObservation) -> str:
    values = [child.gender or "未填写" for child in record.children]
    if values and all(value == "未填写" for value in values):
        return ""
    return _collapse_same(values)


def _indicator_detail(item: ExportIndicator) -> str:
    text = f"{item.code} {item.name}·{LEVEL_LABELS.get(item.level, f'第{item.level}阶')}"
    if item.level_desc:
        text += f"：{item.level_desc}"
    return text


def indicator_text(indicators: List[ExportIndicator]) -> str:
    items = [
        _indicator_detail(item) + (f"（行为证据：{item.evidence}）" if item.evidence else "")
        for item in indicators
    ]
    return f"【关联指标】{'；'.join(items)}" if items else ""


def observation_title(record: ExportObservation) -> str:
    return f"{child_names(record) or '未指定幼儿'}的观察记录"


def observation_date(record: ExportObservation) -> str:
    local = kindergarten_datetime(record.observed_at)
    return f"{local.year}年{local.month}月{local.day}日"


def indicator_lines(indicators: List[ExportIndicator]) -> List[str]:
    lines: List[str] = []
    for item in indicators:
        lines.append(_indicator_detail(item))
        if item.evidence:
            lines.append(f"　行为证据：{item.evidence}")
    return lines


def build_observation_markdown(
    records: List[ExportObservation],
    *,
    include_indicators: bool = False,
) -> bytes:
    """生成适合二次整理的 UTF-8 Markdown；多篇用分隔线分开。"""
    documents = []
    for record in records:
        lines = [
            f"# {observation_title(record)}",
            "",
            f"{observation_date(record)} · {record.area_name}",
            "",
            "## 基本信息",
            "",
            f"- 幼儿：{child_names(record)}",
            f"- 年龄：{child_ages(record)}",
            f"- 性别：{child_genders(record)}",
            f"- 观察者：{record.observer_name}",
            f"- 时间：{observation_date(record)}",
            f"- 场景：{record.location or record.area_name}",
        ]
        if record.background_note:
            lines.append(f"- 背景：{record.background_note}")
        sections = (
            ("观察目的", record.purpose),
            ("观察记录", record.narrative),
            ("观察分析", record.analysis),
        )
        for heading, content in sections:
            lines.extend(["", f"## {heading}", "", content or ""])
        if include_indicators and record.indicators:
            lines.extend(["", "### 已采纳观察指标", ""])
            lines.extend(f"- {item}" for item in indicator_lines(record.indicators))
        lines.extend([
            "",
            "## 下一步支持策略",
            "",
            record.strategy or "",
            "",
            "---",
            "",
            f"记录人：{record.observer_name}",
            "",
            observation_date(record),
        ])
        documents.append("\n".join(lines).rstrip())
    return ("\n\n---\n\n".join(documents) + "\n").encode("utf-8")


def _pdf_paragraph_text(value: Optional[str]) -> str:
    return escape(value or "").replace("\n", "<br/>")


def _pdf_styles():
    try:
        pdfmetrics.getFont(PDF_FONT_NAME)
    except KeyError:
        font_path = next((path for path in PDF_FONT_CANDIDATES if path.is_file()), None)
        if not font_path:
            raise RuntimeError("找不到可嵌入的中文字体，无法生成 PDF")
        pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(font_path), subfontIndex=0))
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "BangbangTitle",
            parent=styles["Title"],
            fontName=PDF_FONT_NAME,
            fontSize=19,
            leading=27,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#225440"),
            spaceAfter=8,
        ),
        "subtitle": ParagraphStyle(
            "BangbangSubtitle",
            parent=styles["Normal"],
            fontName=PDF_FONT_NAME,
            fontSize=10.5,
            leading=16,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#676960"),
            spaceAfter=16,
        ),
        "section": ParagraphStyle(
            "BangbangSection",
            parent=styles["Heading2"],
            fontName=PDF_FONT_NAME,
            fontSize=13,
            leading=19,
            textColor=colors.HexColor("#225440"),
            spaceBefore=10,
            spaceAfter=7,
        ),
        "body": ParagraphStyle(
            "BangbangBody",
            parent=styles["BodyText"],
            fontName=PDF_FONT_NAME,
            fontSize=10.5,
            leading=18,
            alignment=TA_LEFT,
            textColor=colors.HexColor("#242520"),
            spaceAfter=7,
        ),
        "label": ParagraphStyle(
            "BangbangLabel",
            parent=styles["BodyText"],
            fontName=PDF_FONT_NAME,
            fontSize=9.5,
            leading=15,
            textColor=colors.HexColor("#676960"),
        ),
        "signature": ParagraphStyle(
            "BangbangSignature",
            parent=styles["BodyText"],
            fontName=PDF_FONT_NAME,
            fontSize=9.5,
            leading=16,
            alignment=TA_RIGHT,
            textColor=colors.HexColor("#676960"),
        ),
    }


def _pdf_photo_flowables(record: ExportObservation, styles):
    """在 PDF 里把现场照片/视频代表帧放在记录偏前的位置；图片异常会跳过。"""
    out = []
    if not record.highlight_frames:
        return out
    try:
        from PIL import Image as PILImage
        out.append(Paragraph("▸ 现场影像", styles["section"]))
        for frame in record.highlight_frames:
            try:
                img = PILImage.open(BytesIO(frame))
                width = 6 * cm
                height = width * (img.size[1] / img.size[0]) if img.size[0] else width
                out.append(ReportLabImage(BytesIO(frame), width=width, height=height))
            except Exception:
                continue
    except Exception:
        pass
    return out


def _pdf_record_flowables(record: ExportObservation, include_indicators: bool, styles):
    basics = [
        ("幼儿", child_names(record)),
        ("年龄", child_ages(record)),
        ("性别", child_genders(record)),
        ("观察者", record.observer_name),
        ("时间", observation_date(record)),
        ("场景", record.location or record.area_name),
    ]
    if record.background_note:
        basics.append(("背景", record.background_note))
    table_data = [
        [
            Paragraph(f"<b>{_pdf_paragraph_text(label)}</b>", styles["label"]),
            Paragraph(_pdf_paragraph_text(value), styles["body"]),
        ]
        for label, value in basics
    ]
    info_table = Table(table_data, colWidths=[2.2 * cm, 13.6 * cm], hAlign="LEFT")
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F7F6F1")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDD9CF")),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#E5E1D8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story = [
        Paragraph(_pdf_paragraph_text(observation_title(record)), styles["title"]),
        Paragraph(
            _pdf_paragraph_text(f"{observation_date(record)} · {record.area_name}"),
            styles["subtitle"],
        ),
        info_table,
        Spacer(1, 8),
    ]
    story.extend(_pdf_photo_flowables(record, styles))
    for heading, content in (
        ("观察目的", record.purpose),
        ("观察记录", record.narrative),
        ("观察分析", record.analysis),
    ):
        story.extend([
            Paragraph(heading, styles["section"]),
            Paragraph(_pdf_paragraph_text(content), styles["body"]),
        ])
    if include_indicators and record.indicators:
        indicator_items = [
            Paragraph(f"• {_pdf_paragraph_text(item)}", styles["body"])
            for item in indicator_lines(record.indicators)
        ]
        story.append(KeepTogether([
            Paragraph("已采纳观察指标", styles["section"]),
            *indicator_items,
        ]))
    story.extend([
        Paragraph("下一步支持策略", styles["section"]),
        Paragraph(_pdf_paragraph_text(record.strategy), styles["body"]),
        Spacer(1, 12),
        Paragraph(
            (
                f"记录人：{_pdf_paragraph_text(record.observer_name)}"
                f"<br/>{_pdf_paragraph_text(observation_date(record))}"
            ),
            styles["signature"],
        ),
    ])
    return story


def build_observation_pdf(
    records: List[ExportObservation],
    *,
    include_indicators: bool = False,
) -> bytes:
    """生成固定排版 PDF；多篇记录每篇另起一页。"""
    output = BytesIO()
    styles = _pdf_styles()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title="帮帮师记观察记录",
        author="帮帮师记",
    )
    story = []
    for index, record in enumerate(records):
        if index:
            story.append(PageBreak())
        story.extend(_pdf_record_flowables(record, include_indicators, styles))
    document.build(story)
    return output.getvalue()


def _set_font(run, *, size: float = 10.5, bold: bool = False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)


def _set_cell_text(cell, text: str, *, bold: bool = False, center: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text or "")
    _set_font(run, bold=bold)


def _append_cell_paragraph(cell, text: str):
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    _set_font(run)


def _set_cell_margins(cell, *, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def _set_table_geometry(table):
    widths = [Cm(2.5), Cm(6.2), Cm(2.5), Cm(6.2)]
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width.twips)))
        grid.append(grid_col)
    tbl_width = table._tbl.tblPr.first_child_found_in("w:tblW")
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        table._tbl.tblPr.insert(0, tbl_width)
    tbl_width.set(qn("w:w"), str(sum(int(width.twips) for width in widths)))
    tbl_width.set(qn("w:type"), "dxa")
    _set_table_borders(table)


def _add_title(
    document: Document,
    observed_at: datetime,
    *,
    page_break_before: bool = False,
):
    academic_year, term = academic_year_and_term(observed_at)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = True
    if page_break_before:
        title.paragraph_format.page_break_before = True
    title_run = title.add_run("自主游戏观察记录表")
    _set_font(title_run, size=16, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.paragraph_format.keep_with_next = False
    subtitle_run = subtitle.add_run(f"{academic_year}{term}")
    _set_font(subtitle_run, size=12)


def _add_record_table(
    document: Document,
    record: ExportObservation,
    include_indicators: bool,
    *,
    page_break_before: bool = False,
):
    _add_title(document, record.observed_at, page_break_before=page_break_before)
    table = document.add_table(rows=8, cols=4)
    _set_table_geometry(table)

    # 严格按园所模板：第 1、3 列分别纵向合并前三行。
    table.cell(0, 0).merge(table.cell(2, 0))
    table.cell(0, 2).merge(table.cell(2, 2))
    for row_index in range(4, 8):
        table.cell(row_index, 1).merge(table.cell(row_index, 3))

    local = kindergarten_datetime(record.observed_at)
    _set_cell_text(table.cell(0, 0), "观察对象", bold=True, center=True)
    _set_cell_text(table.cell(0, 1), f"姓名：{child_names(record)}")
    _set_cell_text(table.cell(1, 1), f"年龄：{child_ages(record)}")
    _set_cell_text(table.cell(2, 1), f"性别：{child_genders(record)}")
    _set_cell_text(table.cell(0, 2), "观察背景", bold=True, center=True)
    _set_cell_text(table.cell(0, 3), f"观察地点：{record.location or record.area_name}")
    _set_cell_text(table.cell(1, 3), f"其他背景信息：{record.background_note or ''}")
    _set_cell_text(table.cell(2, 3), "其他背景信息：")
    _set_cell_text(table.cell(3, 0), "观察时间", bold=True, center=True)
    _set_cell_text(table.cell(3, 1), f"{local.year}年{local.month}月{local.day}日")
    _set_cell_text(table.cell(3, 2), "观察者", bold=True, center=True)
    _set_cell_text(table.cell(3, 3), record.observer_name)

    body_rows = (
        (4, "观察目的", record.purpose),
        (5, "观察记录", record.narrative),
        (6, "观察分析", record.analysis),
        (7, "下一步支持策略", record.strategy),
    )
    for row_index, label, value in body_rows:
        _set_cell_text(table.cell(row_index, 0), label, bold=True, center=True)
        _set_cell_text(table.cell(row_index, 1), value or "")
        table.cell(row_index, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    if include_indicators:
        appended = indicator_text(record.indicators)
        if appended:
            _append_cell_paragraph(table.cell(6, 1), appended)

    # 把上传照片与视频代表帧放进「观察记录」单元格（表格内，靠后不突兀）。
    # 图片异常时跳过插入，避免坏图让整份导出失败。
    if record.highlight_frames:
        cell = table.cell(5, 1)
        try:
            cap = cell.add_paragraph()
            cap.add_run("▸ 现场影像")
            for frame in record.highlight_frames:
                pic = cell.add_paragraph()
                pic.add_run().add_picture(BytesIO(frame), width=Cm(10))
        except Exception:
            pass


def build_observation_document(
    records: List[ExportObservation],
    *,
    include_indicators: bool = False,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    # 删除 python-docx 默认创建的空段落，避免标题上方多出空白。
    document._body.clear_content()
    for index, record in enumerate(records):
        _add_record_table(
            document,
            record,
            include_indicators,
            page_break_before=index > 0,
        )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
