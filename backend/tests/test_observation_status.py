import json
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import httpx
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import text
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

import main
from export_service import academic_year_and_term
from models import (
    AIRun,
    Area,
    Child,
    ClassRoom,
    Observation,
    ObservationChild,
    ObservationTag,
    Teacher,
)


class ObservationStatusFlowTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.ai_mode_patcher = patch.object(main.ai_service, "AI_MODE", "mock")
        self.ai_mode_patcher.start()
        engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        SQLModel.metadata.create_all(engine)
        main.engine = engine
        main.UPLOAD_DIR = Path(self.temp_dir.name) / "uploads"
        main.UPLOAD_DIR.mkdir()
        self.client = TestClient(main.app)

        with Session(engine) as session:
            area = Area(code="construction", name="建构区")
            room = ClassRoom(name="中二班", age_group="middle")
            session.add(area)
            session.add(room)
            session.commit()
            session.refresh(area)
            session.refresh(room)
            teacher = Teacher(name="测试教师", classroom_id=room.id)
            child = Child(
                name="测试幼儿A",
                classroom_id=room.id,
                birth_date=date(2021, 5, 20),
                gender="女",
            )
            second_child = Child(name="测试幼儿B", classroom_id=room.id, gender="男")
            session.add(teacher)
            session.add(child)
            session.add(second_child)
            session.commit()
            session.refresh(child)
            session.refresh(second_child)
            self.area_id = area.id
            self.child_id = child.id
            self.second_child_id = second_child.id

    def tearDown(self):
        self.ai_mode_patcher.stop()
        self.temp_dir.cleanup()

    def create_bound_observation(self):
        media_response = self.client.post(
            "/uploads?duration_sec=2100",
            files={"file": ("mock.mp4", b"mock-video", "video/mp4")},
        )
        self.assertEqual(media_response.status_code, 201)
        media_id = media_response.json()["id"]

        rejected_status = self.client.post(
            "/observations",
            json={
                "area_id": self.area_id,
                "status": "confirmed",
            },
        )
        self.assertEqual(rejected_status.status_code, 422)

        create_response = self.client.post(
            "/observations",
            json={
                "child_id": self.child_id,
                "area_id": self.area_id,
                "note": "测试状态流转",
            },
        )
        self.assertEqual(create_response.status_code, 201)
        observation = create_response.json()
        self.assertEqual(observation["status"], "uploaded")

        attach_response = self.client.post(
            f"/observations/{observation['id']}/attach-media?media_id={media_id}"
        )
        self.assertEqual(attach_response.status_code, 200)
        self.assertEqual(attach_response.json()["status"], "uploaded")
        return observation["id"]

    def test_complete_flow_filter_and_illegal_transition(self):
        observation_id = self.create_bound_observation()

        illegal = self.client.post(f"/observations/{observation_id}/confirm")
        self.assertEqual(illegal.status_code, 400)
        self.assertEqual(illegal.json()["detail"]["current_status"], "uploaded")
        self.assertEqual(illegal.json()["detail"]["target_status"], "confirmed")

        narrative = self.client.post(f"/observations/{observation_id}/narrative")
        self.assertEqual(narrative.status_code, 200)
        self.assertEqual(narrative.json()["status"], "ready_for_review")
        self.assertIsNotNone(narrative.json()["processing_started_at"])
        self.assertIsNotNone(narrative.json()["ready_at"])
        self.assertTrue(narrative.json()["is_mock"])

        suggestions = self.client.post(f"/observations/{observation_id}/suggest-tags")
        self.assertEqual(suggestions.status_code, 200)
        self.assertEqual(suggestions.json()["status"], "ready_for_review")
        self.assertIsNotNone(suggestions.json()["ready_at"])
        self.assertTrue(suggestions.json()["suggestions"])
        self.assertTrue(suggestions.json()["quant_hits"])
        with Session(main.engine) as session:
            mock_run = session.get(AIRun, suggestions.json()["ai_run_id"])
            self.assertEqual(mock_run.status, "completed")
            self.assertTrue(mock_run.is_mock)
            self.assertEqual(mock_run.provider, "mock")

        detail_with_tags = self.client.get(f"/observations/{observation_id}").json()
        system_tags = [
            tag for tag in detail_with_tags["tags"]
            if tag["source"] == "system_determined"
        ]
        self.assertTrue(system_tags)
        self.assertTrue(all(tag["accepted"] is True for tag in system_tags))

        cancelled_system = self.client.patch(
            f"/observations/{observation_id}/tags/{system_tags[0]['id']}",
            json={"accepted": False},
        )
        self.assertEqual(cancelled_system.status_code, 200)
        restored_detail = self.client.get(f"/observations/{observation_id}").json()
        restored_system = next(
            tag for tag in restored_detail["tags"] if tag["id"] == system_tags[0]["id"]
        )
        self.assertFalse(restored_system["accepted"])

        update_context = self.client.patch(
            f"/observations/{observation_id}",
            json={"child_id": self.child_id, "note": "确认时补充"},
        )
        self.assertEqual(update_context.status_code, 200)
        self.assertEqual(update_context.json()["note"], "确认时补充")

        ready_items = self.client.get(
            "/observations", params={"status": "ready_for_review"}
        )
        self.assertEqual(ready_items.status_code, 200)
        self.assertEqual([item["id"] for item in ready_items.json()], [observation_id])

        first_tag_id = suggestions.json()["suggestions"][0]["tag_id"]
        decision = self.client.patch(
            f"/observations/{observation_id}/tags/{first_tag_id}",
            json={"accepted": True},
        )
        self.assertEqual(decision.status_code, 200)

        second_tag_id = suggestions.json()["suggestions"][1]["tag_id"]
        rejected = self.client.patch(
            f"/observations/{observation_id}/tags/{second_tag_id}",
            json={"accepted": False},
        )
        self.assertEqual(rejected.status_code, 200)

        metrics = self.client.get("/metrics/ai-quality").json()
        self.assertEqual(metrics["AI建议总数"], 3)
        self.assertEqual(metrics["教师已处理"], 2)
        self.assertEqual(metrics["教师采纳"], 1)
        self.assertEqual(metrics["采纳率"], 0.5)

        teacher_tag = self.client.post(
            f"/observations/{observation_id}/tags",
            json={"indicator_code": "4.3", "level": 2},
        )
        self.assertEqual(teacher_tag.status_code, 201)
        self.assertEqual(
            teacher_tag.json()["ai_run_id"],
            suggestions.json()["ai_run_id"],
        )

        confirmed = self.client.post(f"/observations/{observation_id}/confirm")
        self.assertEqual(confirmed.status_code, 200)
        confirmed_observation = confirmed.json()["observation"]
        self.assertEqual(confirmed_observation["status"], "confirmed")
        self.assertIsNotNone(confirmed_observation["confirmed_at"])
        detail = self.client.get(f"/observations/{observation_id}").json()
        self.assertEqual(detail["child_confirmed_count"], 1)

    def test_update_child_and_teacher_basic_information(self):
        child_response = self.client.patch(
            f"/children/{self.child_id}",
            json={
                "name": "幼儿A",
                "birth_date": "2021-08-15",
                "gender": "男",
            },
        )
        self.assertEqual(child_response.status_code, 200)
        self.assertEqual(child_response.json()["name"], "幼儿A")
        self.assertEqual(child_response.json()["birth_date"], "2021-08-15")
        self.assertEqual(child_response.json()["gender"], "男")

        invalid_gender = self.client.patch(
            f"/children/{self.child_id}",
            json={"gender": "未知"},
        )
        self.assertEqual(invalid_gender.status_code, 422)
        unchanged = self.client.get("/children").json()
        saved_child = next(item for item in unchanged if item["id"] == self.child_id)
        self.assertEqual(saved_child["gender"], "男")

        teacher_response = self.client.patch(
            "/teachers/1",
            json={"name": "教师A"},
        )
        self.assertEqual(teacher_response.status_code, 200)
        self.assertEqual(teacher_response.json()["name"], "教师A")
        self.assertEqual(self.client.get("/teachers").json()[0]["name"], "教师A")

    def test_confirmation_requires_child_without_changing_age_snapshot(self):
        created = self.client.post("/observations", json={"area_id": self.area_id}).json()
        media = self.client.post(
            "/uploads",
            files={"file": ("mock.mp4", b"mock-video", "video/mp4")},
        ).json()
        self.client.post(
            f"/observations/{created['id']}/attach-media",
            params={"media_id": media["id"]},
        )
        self.client.post(f"/observations/{created['id']}/narrative")
        suggestions = self.client.post(
            f"/observations/{created['id']}/suggest-tags"
        ).json()
        self.client.patch(
            f"/observations/{created['id']}/tags/{suggestions['suggestions'][0]['tag_id']}",
            json={"accepted": True},
        )

        missing_child = self.client.post(f"/observations/{created['id']}/confirm")
        self.assertEqual(missing_child.status_code, 400)
        self.assertEqual(
            missing_child.json()["detail"],
            "请先选择这条记录关于哪位幼儿",
        )

        original_classroom_id = created["classroom_id"]
        original_age_group = created["age_group"]
        patched = self.client.patch(
            f"/observations/{created['id']}",
            json={"child_id": self.child_id},
        ).json()
        self.assertEqual(patched["classroom_id"], original_classroom_id)
        self.assertEqual(patched["age_group"], original_age_group)

        confirmed = self.client.post(f"/observations/{created['id']}/confirm")
        self.assertEqual(confirmed.status_code, 200)
        self.assertIsNotNone(confirmed.json()["observation"]["confirmed_at"])

    def test_ai_failure_and_retry(self):
        observation_id = self.create_bound_observation()

        with patch.object(
            main.ai_service,
            "generate_narrative",
            side_effect=RuntimeError("mock workflow unavailable"),
        ):
            failed = self.client.post(f"/observations/{observation_id}/narrative")

        self.assertEqual(failed.status_code, 500)
        with Session(main.engine) as session:
            observation = session.get(Observation, observation_id)
            self.assertEqual(observation.status, "failed")
            self.assertIn("mock workflow unavailable", observation.failure_reason)

        retried = self.client.post(f"/observations/{observation_id}/narrative")
        self.assertEqual(retried.status_code, 200)
        self.assertEqual(retried.json()["status"], "ready_for_review")
        self.assertIsNotNone(retried.json()["ready_at"])

    def test_processing_status_is_visible_while_ai_is_running(self):
        observation_id = self.create_bound_observation()
        started = threading.Event()
        release = threading.Event()
        original_generate = main.ai_service.generate_narrative

        def slow_generate(*args, **kwargs):
            started.set()
            self.assertTrue(release.wait(timeout=2))
            return original_generate(*args, **kwargs)

        with patch.object(main.ai_service, "generate_narrative", side_effect=slow_generate):
            with ThreadPoolExecutor(max_workers=1) as executor:
                future = executor.submit(
                    self.client.post,
                    f"/observations/{observation_id}/narrative",
                )
                self.assertTrue(started.wait(timeout=2))
                processing = self.client.get(f"/observations/{observation_id}")
                self.assertEqual(processing.json()["status"], "processing")
                self.assertIsNotNone(processing.json()["processing_started_at"])
                self.assertIsNone(processing.json()["ready_at"])
                release.set()
                completed = future.result(timeout=2)

        self.assertEqual(completed.json()["status"], "ready_for_review")
        self.assertIsNotNone(completed.json()["ready_at"])

    def test_quick_capture_defaults_media_inference_and_child_update(self):
        created = self.client.post("/observations", json={"area_id": self.area_id})
        self.assertEqual(created.status_code, 201)
        observation = created.json()
        self.assertEqual(observation["status"], "uploaded")
        self.assertIsNone(observation["child_id"])
        self.assertEqual(observation["classroom_id"], 1)
        self.assertEqual(observation["observer_id"], 1)
        self.assertEqual(observation["age_group"], "middle")
        self.assertIsNone(observation["media_type"])

        missing_area = self.client.post("/observations", json={})
        self.assertEqual(missing_area.status_code, 422)

        jpg = self.client.post(
            "/uploads",
            files={"file": ("mock.jpg", b"mock-image", "image/jpeg")},
        )
        self.assertEqual(jpg.status_code, 201)
        attached_jpg = self.client.post(
            f"/observations/{observation['id']}/attach-media",
            params={"media_id": jpg.json()["id"]},
        )
        self.assertEqual(attached_jpg.status_code, 200)
        self.assertEqual(attached_jpg.json()["media_type"], "image")

        second_media = self.client.post(
            "/uploads",
            files={"file": ("second.mp4", b"second-video", "video/mp4")},
        )
        attached_second = self.client.post(
            f"/observations/{observation['id']}/attach-media",
            params={"media_id": second_media.json()["id"]},
        )
        self.assertEqual(attached_second.status_code, 200)
        self.assertEqual(attached_second.json()["media_type"], "image")

        video_observation = self.client.post(
            "/observations", json={"area_id": self.area_id}
        ).json()
        mp4 = self.client.post(
            "/uploads",
            files={"file": ("mock.mp4", b"mock-video", "video/mp4")},
        )
        self.assertEqual(mp4.status_code, 201)
        attached_mp4 = self.client.post(
            f"/observations/{video_observation['id']}/attach-media",
            params={"media_id": mp4.json()["id"]},
        )
        self.assertEqual(attached_mp4.status_code, 200)
        self.assertEqual(attached_mp4.json()["media_type"], "video")

        original_classroom_id = observation["classroom_id"]
        original_age_group = observation["age_group"]
        updated = self.client.patch(
            f"/observations/{observation['id']}",
            json={"child_id": self.child_id, "note": "回看时确认幼儿"},
        )
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.json()["child_id"], self.child_id)
        self.assertEqual(updated.json()["note"], "回看时确认幼儿")
        self.assertEqual(updated.json()["classroom_id"], original_classroom_id)
        self.assertEqual(updated.json()["age_group"], original_age_group)
        detail = self.client.get(f"/observations/{observation['id']}").json()
        self.assertEqual(detail["observer"]["name"], "测试教师")
        self.assertEqual(len(detail["children"]), 1)
        self.assertEqual(detail["children"][0]["id"], self.child_id)
        self.assertTrue(detail["children"][0]["is_primary"])
        self.assertEqual(detail["children"][0]["birth_date"], "2021-05-20")
        self.assertEqual(detail["children"][0]["gender"], "女")

    def test_multi_child_detail_and_confirmed_count_use_association(self):
        with Session(main.engine) as session:
            first = Observation(
                child_id=self.child_id,
                area_id=self.area_id,
                classroom_id=1,
                observer_id=1,
                age_group="middle",
                status="confirmed",
            )
            second = Observation(
                child_id=self.second_child_id,
                area_id=self.area_id,
                classroom_id=1,
                observer_id=1,
                age_group="middle",
                status="confirmed",
            )
            session.add(first)
            session.add(second)
            session.flush()
            session.add_all([
                ObservationChild(
                    observation_id=first.id,
                    child_id=self.child_id,
                    is_primary=True,
                ),
                ObservationChild(
                    observation_id=first.id,
                    child_id=self.second_child_id,
                    is_primary=False,
                ),
                ObservationChild(
                    observation_id=second.id,
                    child_id=self.second_child_id,
                    is_primary=True,
                ),
            ])
            session.commit()
            first_id = first.id

        detail = self.client.get(f"/observations/{first_id}")
        self.assertEqual(detail.status_code, 200)
        body = detail.json()
        self.assertEqual(body["child_confirmed_count"], 1)
        self.assertEqual(
            [(item["id"], item["is_primary"]) for item in body["children"]],
            [(self.child_id, True), (self.second_child_id, False)],
        )
        counts = {
            item["id"]: item["confirmed_observation_count"]
            for item in body["children"]
        }
        self.assertEqual(counts[self.child_id], 1)
        self.assertEqual(counts[self.second_child_id], 2)

    def test_docx_export_structure_fields_indicators_and_monthly(self):
        observed_at = datetime(2026, 8, 23, 17, 57, tzinfo=timezone.utc)
        with Session(main.engine) as session:
            observation = Observation(
                child_id=self.child_id,
                area_id=self.area_id,
                classroom_id=1,
                observer_id=1,
                observed_at=observed_at,
                age_group="middle",
                status="confirmed",
                location="楼顶建构区",
                background_note="幼儿连续搭建第三天",
                purpose="观察幼儿解决问题的过程",
                narrative="1. 先放下长条积木。\n2. 调整间距后继续搭建。",
                analysis="幼儿会根据结果调整计划。",
                strategy="提供更多不同形状的材料。",
            )
            session.add(observation)
            session.flush()
            session.add_all([
                ObservationChild(
                    observation_id=observation.id,
                    child_id=self.child_id,
                    is_primary=True,
                ),
                ObservationChild(
                    observation_id=observation.id,
                    child_id=self.second_child_id,
                    is_primary=False,
                ),
                ObservationTag(
                    observation_id=observation.id,
                    indicator_code="4.4",
                    indicator_name="试误与问题解决",
                    level=2,
                    source="ai_suggested",
                    accepted=True,
                ),
                ObservationTag(
                    observation_id=observation.id,
                    indicator_code="3.1",
                    indicator_name="互动形式",
                    level=3,
                    source="ai_suggested",
                    accepted=False,
                ),
                ObservationTag(
                    observation_id=observation.id,
                    indicator_code="1.2",
                    indicator_name="身体探索方式",
                    level=1,
                    source="system_determined",
                    accepted=True,
                ),
            ])
            second_observation = Observation(
                child_id=self.child_id,
                area_id=self.area_id,
                classroom_id=1,
                observer_id=1,
                observed_at=datetime(2026, 8, 24, 2, 0, tzinfo=timezone.utc),
                age_group="middle",
                status="confirmed",
                narrative="第二条月度记录",
            )
            session.add(second_observation)
            session.flush()
            session.add(ObservationChild(
                observation_id=second_observation.id,
                child_id=self.child_id,
                is_primary=True,
            ))
            session.commit()
            observation_id = observation.id

        exported = self.client.get(f"/observations/{observation_id}/export")
        self.assertEqual(exported.status_code, 200)
        self.assertEqual(
            exported.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn(
            "%E8%A7%82%E5%AF%9F%E8%AE%B0%E5%BD%95_",
            exported.headers["content-disposition"],
        )
        document = Document(BytesIO(exported.content))
        self.assertEqual(len(document.tables), 1)
        table = document.tables[0]
        self.assertEqual(len(table.rows), 8)
        self.assertEqual(len(table.columns), 4)
        self.assertIs(table.cell(0, 0)._tc, table.cell(2, 0)._tc)
        self.assertIs(table.cell(0, 2)._tc, table.cell(2, 2)._tc)
        self.assertIs(table.cell(4, 1)._tc, table.cell(4, 3)._tc)
        all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        all_text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
        self.assertIn("自主游戏观察记录表", all_text)
        self.assertIn("2025-2026学年度第二学期", all_text)
        self.assertIn("姓名：测试幼儿A、测试幼儿B", all_text)
        self.assertIn("年龄：5岁、中班", all_text)
        self.assertIn("性别：女、男", all_text)
        self.assertIn("观察地点：楼顶建构区", all_text)
        self.assertIn("2026年8月24日", all_text)
        self.assertIn("测试教师", all_text)
        self.assertNotIn("【关联指标】", all_text)

        with_indicators = self.client.get(
            f"/observations/{observation_id}/export",
            params={"include_indicators": "true"},
        )
        indicator_document = Document(BytesIO(with_indicators.content))
        analysis_text = indicator_document.tables[0].cell(6, 1).text
        self.assertIn("【关联指标】4.4 试误与问题解决·中阶；1.2 身体探索方式·初阶", analysis_text)
        self.assertNotIn("3.1 互动形式", analysis_text)

        markdown = self.client.get(
            f"/observations/{observation_id}/export",
            params={"include_indicators": "true", "format": "md"},
        )
        self.assertEqual(markdown.status_code, 200)
        self.assertEqual(markdown.headers["content-type"], "text/markdown; charset=utf-8")
        self.assertIn(".md", markdown.headers["content-disposition"])
        markdown_text = markdown.content.decode("utf-8")
        for expected in (
            "# 测试幼儿A、测试幼儿B的观察记录",
            "## 基本信息",
            "幼儿：测试幼儿A、测试幼儿B",
            "观察者：测试教师",
            "## 观察目的",
            "观察幼儿解决问题的过程",
            "## 观察记录",
            "1. 先放下长条积木。",
            "## 观察分析",
            "4.4 试误与问题解决·中阶",
            "1.2 身体探索方式·初阶",
            "## 下一步支持策略",
            "提供更多不同形状的材料。",
            "记录人：测试教师",
        ):
            self.assertIn(expected, markdown_text)
        self.assertNotIn("3.1 互动形式", markdown_text)

        pdf = self.client.get(
            f"/observations/{observation_id}/export",
            params={"include_indicators": "true", "format": "pdf"},
        )
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.headers["content-type"], "application/pdf")
        self.assertIn(".pdf", pdf.headers["content-disposition"])
        self.assertTrue(pdf.content.startswith(b"%PDF-"))

        invalid_format = self.client.get(
            f"/observations/{observation_id}/export",
            params={"format": "txt"},
        )
        self.assertEqual(invalid_format.status_code, 422)
        self.assertEqual(
            invalid_format.json()["detail"],
            "导出格式只支持 docx、pdf 或 md",
        )

        monthly = self.client.get("/exports/monthly", params={"year": 2026, "month": 8})
        self.assertEqual(monthly.status_code, 200)
        monthly_document = Document(BytesIO(monthly.content))
        self.assertEqual(len(monthly_document.tables), 2)
        page_breaks = monthly_document._element.xpath('.//w:pageBreakBefore')
        self.assertEqual(len(page_breaks), 1)

        monthly_pdf = self.client.get(
            "/exports/monthly",
            params={"year": 2026, "month": 8, "format": "pdf"},
        )
        self.assertEqual(monthly_pdf.status_code, 200)
        self.assertTrue(monthly_pdf.content.startswith(b"%PDF-"))
        monthly_markdown = self.client.get(
            "/exports/monthly",
            params={"year": 2026, "month": 8, "format": "md"},
        )
        self.assertEqual(monthly_markdown.status_code, 200)
        markdown_titles = [
            line for line in monthly_markdown.content.decode("utf-8").splitlines()
            if line.startswith("# ")
        ]
        self.assertEqual(len(markdown_titles), 2)

        missing_month = self.client.get(
            "/exports/monthly", params={"year": 2025, "month": 7}
        )
        self.assertEqual(missing_month.status_code, 404)
        self.assertEqual(
            missing_month.json()["detail"],
            "2025年7月没有已确认的观察记录",
        )

        with Session(main.engine) as session:
            unconfirmed = Observation(
                area_id=self.area_id,
                classroom_id=1,
                observer_id=1,
                age_group="middle",
                status="ready_for_review",
            )
            session.add(unconfirmed)
            session.commit()
            unconfirmed_id = unconfirmed.id
        rejected = self.client.get(f"/observations/{unconfirmed_id}/export")
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(rejected.json()["detail"], "未确认的记录不能导出")

    def test_academic_year_term_boundaries(self):
        self.assertEqual(
            academic_year_and_term(datetime(2025, 1, 10, tzinfo=timezone.utc)),
            ("2024-2025学年度", "第一学期"),
        )
        self.assertEqual(
            academic_year_and_term(datetime(2025, 2, 10, tzinfo=timezone.utc)),
            ("2024-2025学年度", "第二学期"),
        )
        self.assertEqual(
            academic_year_and_term(datetime(2025, 8, 10, tzinfo=timezone.utc)),
            ("2024-2025学年度", "第二学期"),
        )
        self.assertEqual(
            academic_year_and_term(datetime(2025, 9, 10, tzinfo=timezone.utc)),
            ("2025-2026学年度", "第一学期"),
        )

    def test_upload_limit_and_read_media_file(self):
        image_content = b"mock-image-content"
        uploaded = self.client.post(
            "/uploads",
            files={"file": ("mock.jpg", image_content, "image/jpeg")},
        )
        self.assertEqual(uploaded.status_code, 201)

        media_file = self.client.get(f"/media/{uploaded.json()['id']}/file")
        self.assertEqual(media_file.status_code, 200)
        self.assertEqual(media_file.headers["content-type"], "image/jpeg")
        self.assertEqual(media_file.content, image_content)

        missing = self.client.get("/media/999999/file")
        self.assertEqual(missing.status_code, 404)

        with patch.object(main, "MAX_SIZE", 8):
            too_large = self.client.post(
                "/uploads",
                files={"file": ("large.mp4", b"123456789", "video/mp4")},
            )
        self.assertEqual(too_large.status_code, 413)
        self.assertEqual(too_large.json()["detail"], "文件不能超过 200MB")

    def test_upload_rejects_video_over_duration_limit(self):
        with patch.object(main, "_probe_video_duration", return_value=300):
            resp = self.client.post(
                "/uploads",
                files={"file": ("long.mp4", b"mock-video", "video/mp4")},
            )
        self.assertEqual(resp.status_code, 422)
        self.assertIn("3 分钟", resp.json()["detail"])

    def test_upload_stores_real_video_duration(self):
        with patch.object(main, "_probe_video_duration", return_value=120):
            resp = self.client.post(
                "/uploads",
                files={"file": ("clip.mp4", b"mock-video", "video/mp4")},
            )
        self.assertEqual(resp.status_code, 201)
        self.assertEqual(resp.json()["duration_sec"], 120)

    def test_iphone_media_types_and_extension_fallback(self):
        cases = [
            ("iphone.mov", "video/quicktime", "video/quicktime", "video"),
            ("iphone.heic", "image/heic", "image/heic", "image"),
            ("iphone.heif", "image/heif", "image/heif", "image"),
            ("fallback.mov", "application/octet-stream", "video/quicktime", "video"),
            ("fallback.heic", "", "image/heic", "image"),
            ("fallback.heif", "application/octet-stream", "image/heif", "image"),
            ("fallback.jpg", "application/octet-stream", "image/jpeg", "image"),
            ("fallback.jpeg", "application/octet-stream", "image/jpeg", "image"),
            ("fallback.png", "application/octet-stream", "image/png", "image"),
            ("fallback.mp4", "application/octet-stream", "video/mp4", "video"),
        ]

        for filename, sent_type, expected_type, expected_media_type in cases:
            with self.subTest(filename=filename, sent_type=sent_type):
                uploaded = self.client.post(
                    "/uploads",
                    files={"file": (filename, b"mock-iphone-media", sent_type)},
                )
                self.assertEqual(uploaded.status_code, 201)
                self.assertEqual(uploaded.json()["content_type"], expected_type)

                observation = self.client.post(
                    "/observations", json={"area_id": self.area_id}
                ).json()
                attached = self.client.post(
                    f"/observations/{observation['id']}/attach-media",
                    params={"media_id": uploaded.json()["id"]},
                )
                self.assertEqual(attached.status_code, 200)
                self.assertEqual(attached.json()["media_type"], expected_media_type)

        rejected = self.client.post(
            "/uploads",
            files={"file": ("notes.txt", b"not-media", "text/plain")},
        )
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(
            rejected.json()["detail"],
            "只支持照片和视频（JPG、PNG、HEIC、MP4、MOV）",
        )

    def test_utc_timestamps_and_legacy_local_time_response(self):
        created = self.client.post("/observations", json={"area_id": self.area_id})
        self.assertEqual(created.status_code, 201)
        created_at = created.json()["created_at"]
        self.assertTrue(created_at.endswith("Z") or created_at.endswith("+00:00"))

        with main.engine.begin() as connection:
            stored = connection.execute(
                text("SELECT created_at FROM observation WHERE id = :id"),
                {"id": created.json()["id"]},
            ).scalar_one()
            self.assertTrue(stored.endswith("Z") or stored.endswith("+00:00"))
            connection.execute(
                text("UPDATE observation SET created_at = :created_at WHERE id = :id"),
                {
                    "id": created.json()["id"],
                    "created_at": "2026-08-23 05:34:32.426166",
                },
            )

        legacy = self.client.get(f"/observations/{created.json()['id']}")
        self.assertEqual(legacy.status_code, 200)
        expected = datetime(2026, 8, 23, 12, 34, 32, 426166, tzinfo=timezone.utc)
        actual = datetime.fromisoformat(legacy.json()["created_at"].replace("Z", "+00:00"))
        self.assertEqual(actual, expected)

    def test_video_thumbnail_success_and_failure_reason(self):
        def create_thumbnail(source_path):
            main.get_thumbnail_path(source_path.name).write_bytes(b"mock-png")
            return None

        with patch.object(main, "generate_video_thumbnail", side_effect=create_thumbnail):
            uploaded = self.client.post(
                "/uploads",
                files={"file": ("iphone.mov", b"mock-mov", "video/quicktime")},
            )

        self.assertEqual(uploaded.status_code, 201)
        thumbnail = self.client.get(f"/media/{uploaded.json()['id']}/thumbnail")
        self.assertEqual(thumbnail.status_code, 200)
        self.assertEqual(thumbnail.headers["content-type"], "image/png")
        self.assertEqual(thumbnail.content, b"mock-png")

        def fail_thumbnail(source_path):
            reason = "当前视频编码暂时无法生成缩略图"
            main.get_thumbnail_error_path(source_path.name).write_text(reason, encoding="utf-8")
            return reason

        with patch.object(main, "generate_video_thumbnail", side_effect=fail_thumbnail):
            failed_upload = self.client.post(
                "/uploads",
                files={"file": ("unsupported.mov", b"bad-codec", "video/quicktime")},
            )
        failed_media = failed_upload.json()
        self.assertIsNotNone(failed_media["thumbnail_failure_reason"])
        failed_thumbnail = self.client.get(f"/media/{failed_media['id']}/thumbnail")
        self.assertEqual(failed_thumbnail.status_code, 404)
        self.assertEqual(
            failed_thumbnail.json()["detail"]["reason"],
            failed_media["thumbnail_failure_reason"],
        )

        with patch.object(main, "generate_video_thumbnail", side_effect=create_thumbnail):
            retried_thumbnail = self.client.get(f"/media/{failed_media['id']}/thumbnail")
        self.assertEqual(retried_thumbnail.status_code, 200)
        refreshed_media = self.client.get("/media").json()
        retried_media = next(item for item in refreshed_media if item["id"] == failed_media["id"])
        self.assertIsNone(retried_media["thumbnail_failure_reason"])

    def test_deepseek_suggestions_are_validated_audited_and_linked(self):
        self.assertEqual(
            main.ai_service._anonymize_narrative(
                "张小雨和小雨一起搭积木，李在旁边看",
                ["小雨", "张小雨", "李"],
            ),
            "幼儿B和幼儿A一起搭积木，李在旁边看",
        )
        observation_id = self.create_bound_observation()
        self.client.post(f"/observations/{observation_id}/narrative")
        original_narrative = "张小雨坐在地垫上，李小明调整了间距后继续摆放"
        with Session(main.engine) as session:
            first_child = session.get(Child, self.child_id)
            first_child.name = "李小明"
            second_child = session.get(Child, self.second_child_id)
            second_child.name = "张小雨"
            old_primary = session.get(ObservationChild, (observation_id, self.child_id))
            old_primary.is_primary = False
            session.add(ObservationChild(
                observation_id=observation_id,
                child_id=self.second_child_id,
                is_primary=True,
            ))
            observation = session.get(Observation, observation_id)
            observation.child_id = self.second_child_id
            observation.narrative = original_narrative
            session.add(first_child)
            session.add(second_child)
            session.add(old_primary)
            session.add(observation)
            session.commit()
        model_content = {
            "suggestions": [
                {
                    "indicator_code": "4.4",
                    "indicator_name": "模型不能决定名称",
                    "level": 2,
                    "level_desc": "模型不能决定描述",
                    "confidence": 0.86,
                    "reason": "白描原文写道：“幼儿A坐在地垫上”",
                    "evidence_based": True,
                    "rank": 9,
                },
                {
                    "indicator_code": "1.2",
                    "indicator_name": "身体探索方式",
                    "level": 3,
                    "level_desc": "无效描述",
                    "confidence": 0.8,
                    "reason": "白描原文写道：“白描中不存在的片段”",
                    "evidence_based": True,
                    "rank": 1,
                },
                {
                    "indicator_code": "1.1",
                    "indicator_name": "身体行为参与度",
                    "level": 2,
                    "level_desc": "模型越过系统判定边界",
                    "confidence": 0.88,
                    "reason": "白描原文写道：“幼儿A坐在地垫上”",
                    "evidence_based": True,
                    "rank": 2,
                },
                {
                    "indicator_code": "2.1",
                    "indicator_name": "伦理禁区",
                    "level": 1,
                    "level_desc": "不应出现",
                    "confidence": 0.9,
                    "reason": "白描原文写道：“幼儿A坐在地垫上”",
                    "evidence_based": True,
                    "rank": 2,
                },
            ]
        }
        raw_response = {
            "id": "test-completion",
            "choices": [{
                "message": {"content": json.dumps(model_content, ensure_ascii=False)}
            }],
            "usage": {
                "prompt_tokens": 1200,
                "completion_tokens": 180,
                "total_tokens": 1380,
            },
        }
        response = httpx.Response(
            200,
            json=raw_response,
            request=httpx.Request("POST", "https://api.deepseek.com/chat/completions"),
        )

        with (
            patch.object(main.ai_service, "AI_MODE", "deepseek"),
            patch.object(main.ai_service, "DEEPSEEK_API_KEY", "test-key"),
            patch.object(main.ai_service.httpx, "post", return_value=response) as post,
        ):
            generated = self.client.post(
                f"/observations/{observation_id}/suggest-tags"
            )

        self.assertEqual(generated.status_code, 200)
        body = generated.json()
        self.assertFalse(body["is_mock"])
        self.assertEqual([item["indicator_code"] for item in body["suggestions"]], ["4.4", "1.2"])
        self.assertNotIn("1.1", [item["indicator_code"] for item in body["suggestions"]])
        self.assertIn(
            ("1.1", 2),
            {(item["indicator_code"], item["level"]) for item in body["quant_hits"]},
        )
        self.assertTrue(body["suggestions"][0]["evidence_based"])
        self.assertEqual(
            body["suggestions"][0]["reason"],
            "白描原文：“幼儿A坐在地垫上”",
        )
        self.assertEqual(body["suggestions"][0]["confidence"], 0.86)
        self.assertFalse(body["suggestions"][1]["evidence_based"])
        self.assertEqual(body["suggestions"][1]["confidence"], 0.42)

        request_body = post.call_args.kwargs["json"]
        self.assertEqual(request_body["model"], "deepseek-chat")
        self.assertEqual(request_body["temperature"], 1.0)
        self.assertEqual(request_body["response_format"], {"type": "json_object"})
        rendered = request_body["messages"][1]["content"]
        self.assertNotIn("张小雨", rendered)
        self.assertNotIn("李小明", rendered)
        self.assertIn("幼儿A坐在地垫上", rendered)
        self.assertIn("幼儿B调整了间距后继续摆放", rendered)
        self.assertIn('"indicator_code": "4.4"', rendered)
        self.assertNotIn('"indicator_code": "2.1"', rendered)
        self.assertIn("由系统按客观数据计算，不在你的判定范围内", rendered)

        catalog, excluded = main.ai_service._model_indicator_catalog()
        indicator_11 = next(
            item for item in catalog if item["indicator_code"] == "1.1"
        )
        self.assertEqual([level["level"] for level in indicator_11["levels"]], [3])
        self.assertEqual(
            {
                (item["indicator_code"], item["level"])
                for item in excluded
                if item["indicator_code"] == "1.1"
            },
            {("1.1", 1), ("1.1", 2)},
        )

        with Session(main.engine) as session:
            run = session.get(AIRun, body["ai_run_id"])
            self.assertEqual(run.status, "completed")
            self.assertEqual(run.prompt_version, "wf-b-v2")
            self.assertEqual(run.temperature, 1.0)
            self.assertEqual(run.token_usage["total_tokens"], 1380)
            self.assertNotIn("张小雨", run.prompt_rendered)
            self.assertNotIn("李小明", run.prompt_rendered)
            self.assertIn("幼儿A坐在地垫上", run.prompt_rendered)
            self.assertIn("幼儿B调整了间距后继续摆放", run.prompt_rendered)
            self.assertIn("2.1", run.error_reason)
            self.assertIn("quant_rule", run.error_reason)
            self.assertIn("原文引用无法", run.error_reason)
            tags = session.exec(
                select(ObservationTag).where(
                    ObservationTag.observation_id == observation_id
                )
            ).all()
            self.assertTrue(tags)
            self.assertTrue(all(tag.ai_run_id == run.id for tag in tags))
            saved_observation = session.get(Observation, observation_id)
            self.assertEqual(saved_observation.narrative, original_narrative)

    def test_deepseek_auth_failure_retries_once_and_falls_back(self):
        observation_id = self.create_bound_observation()
        self.client.post(f"/observations/{observation_id}/narrative")
        unauthorized = httpx.Response(
            401,
            json={"error": {"message": "Authentication Fails"}},
            request=httpx.Request("POST", "https://api.deepseek.com/chat/completions"),
        )

        with (
            patch.object(main.ai_service, "AI_MODE", "deepseek"),
            patch.object(main.ai_service, "DEEPSEEK_API_KEY", "wrong-key"),
            patch.object(
                main.ai_service.httpx,
                "post",
                return_value=unauthorized,
            ) as post,
        ):
            generated = self.client.post(
                f"/observations/{observation_id}/suggest-tags"
            )

        self.assertEqual(generated.status_code, 200)
        body = generated.json()
        self.assertTrue(body["is_mock"])
        self.assertEqual(post.call_count, 2)
        self.assertTrue(body["suggestions"])

        with Session(main.engine) as session:
            run = session.get(AIRun, body["ai_run_id"])
            self.assertEqual(run.status, "failed")
            self.assertTrue(run.is_mock)
            self.assertIn("401", run.error_reason)
            tags = session.exec(
                select(ObservationTag).where(
                    ObservationTag.observation_id == observation_id
                )
            ).all()
            self.assertTrue(all(tag.ai_run_id == run.id for tag in tags))

    def test_ai_quality_is_grouped_cleanly_by_prompt_version(self):
        legacy_observation = self.client.post(
            "/observations", json={"area_id": self.area_id}
        ).json()
        v2_observation = self.client.post(
            "/observations", json={"area_id": self.area_id}
        ).json()

        with Session(main.engine) as session:
            completed_run = AIRun(
                observation_id=v2_observation["id"],
                workflow="indicator_suggestion",
                provider="deepseek",
                model="deepseek-chat",
                prompt_version="wf-b-v2",
                status="completed",
                latency_ms=1000,
                response_raw={},
                is_mock=False,
                prompt_rendered="test prompt",
                token_usage={"total_tokens": 400},
                temperature=1.0,
            )
            failed_run = AIRun(
                observation_id=v2_observation["id"],
                workflow="indicator_suggestion",
                provider="deepseek",
                model="deepseek-chat",
                prompt_version="wf-b-v2",
                status="failed",
                latency_ms=3000,
                response_raw={},
                error_reason="test failure",
                is_mock=True,
                prompt_rendered="test prompt",
                temperature=1.0,
            )
            session.add(completed_run)
            session.add(failed_run)
            session.commit()
            session.refresh(completed_run)

            session.add_all([
                ObservationTag(
                    observation_id=legacy_observation["id"],
                    indicator_code="1.3",
                    indicator_name="身体行为复杂性",
                    level=2,
                    source="ai_suggested",
                    accepted=True,
                ),
                ObservationTag(
                    observation_id=legacy_observation["id"],
                    indicator_code="4.4",
                    indicator_name="试误与问题解决",
                    level=2,
                    source="ai_suggested",
                    accepted=False,
                ),
                ObservationTag(
                    observation_id=legacy_observation["id"],
                    indicator_code="1.2",
                    indicator_name="身体探索方式",
                    level=2,
                    source="ai_suggested",
                    accepted=None,
                ),
                ObservationTag(
                    observation_id=legacy_observation["id"],
                    indicator_code="4.3",
                    indicator_name="分析与规划",
                    level=2,
                    source="teacher_added",
                    accepted=True,
                ),
                ObservationTag(
                    observation_id=v2_observation["id"],
                    ai_run_id=completed_run.id,
                    indicator_code="1.3",
                    indicator_name="身体行为复杂性",
                    level=2,
                    source="ai_suggested",
                    accepted=True,
                ),
                ObservationTag(
                    observation_id=v2_observation["id"],
                    ai_run_id=completed_run.id,
                    indicator_code="4.4",
                    indicator_name="试误与问题解决",
                    level=2,
                    source="ai_suggested",
                    accepted=False,
                ),
                ObservationTag(
                    observation_id=v2_observation["id"],
                    ai_run_id=completed_run.id,
                    indicator_code="4.3",
                    indicator_name="分析与规划",
                    level=2,
                    source="teacher_added",
                    accepted=True,
                ),
            ])
            session.commit()

        metrics = self.client.get("/metrics/ai-quality")
        self.assertEqual(metrics.status_code, 200)
        body = metrics.json()
        self.assertEqual(body["AI建议总数"], 5)
        self.assertEqual(body["教师已处理"], 4)
        self.assertEqual(body["教师采纳"], 2)
        self.assertEqual(body["采纳率"], 0.5)
        self.assertEqual(body["教师自己补的"], 2)
        self.assertEqual(body["漏检率"], 0.5)

        legacy = body["按prompt_version"]["rule-mock-v1"]
        self.assertEqual(legacy["provider"], "mock")
        self.assertEqual(legacy["AI建议总数"], 3)
        self.assertEqual(legacy["教师已处理"], 2)
        self.assertEqual(legacy["教师采纳"], 1)
        self.assertEqual(legacy["教师自己补的"], 1)
        self.assertEqual(legacy["调用统计"]["总调用次数"], 0)

        v2 = body["按prompt_version"]["wf-b-v2"]
        self.assertEqual(v2["provider"], "deepseek")
        self.assertEqual(v2["model"], "deepseek-chat")
        self.assertEqual(v2["temperature"], 1.0)
        self.assertEqual(v2["AI建议总数"], 2)
        self.assertEqual(v2["教师已处理"], 2)
        self.assertEqual(v2["教师采纳"], 1)
        self.assertEqual(v2["教师自己补的"], 1)
        self.assertEqual(v2["调用统计"]["平均latency_ms"], 2000.0)
        self.assertEqual(v2["调用统计"]["平均token消耗"], 400.0)
        self.assertEqual(v2["调用统计"]["总调用次数"], 2)
        self.assertEqual(v2["调用统计"]["失败次数"], 1)


if __name__ == "__main__":
    unittest.main()
